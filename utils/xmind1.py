"""
XMind 测试用例转 Excel 工具
运行方式：
    1. 安装依赖：pip install streamlit xmindparser pandas openpyxl pillow
    2. 运行：streamlit run app.py
"""

import io
import json
import os
import re
import tempfile
import zipfile

import pandas as pd
import streamlit as st
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PILImage
from xmindparser import xmind_to_dict


# ============================================================
# 页面配置
# ============================================================
st.set_page_config(
    page_title="XMind 测试用例转 Excel",
    page_icon="📊",
    layout="wide",
)

st.title("📊 XMind 测试用例转 Excel")
st.markdown("上传 XMind 文件，自动解析并导出为 Excel 测试用例表。")


# ============================================================
# 常量
# ============================================================
COLUMNS = ["模块名称", "用例名称", "操作步骤", "预期结果", "实际结果"]

LEVEL_MAP = {
    0: "模块名称",
    1: "用例名称",
    2: "操作步骤",
}
OPERATION_LEVEL = 3

PREPARE_KEYWORD = "测试准备"
DETAIL_KEYWORD = "需求详情"
DETAIL_SHEET_NAME = "需求详情"

IMG_MAX_WIDTH = 200
IMG_COL_WIDTH = 28
IMG_ROW_HEIGHT = 100
IMG_CELL_PADDING = 4


# ============================================================
# Sheet 名处理
# ============================================================
def _safe_sheet_name(name, default="测试用例"):
    if not name:
        return default
    cleaned = re.sub(r"[:\\/?*\[\]]", "_", str(name)).strip()
    cleaned = cleaned.strip("'")
    if not cleaned:
        cleaned = default
    return cleaned[:31]


# ============================================================
# 通用工具
# ============================================================
def _children_of(topic):
    if isinstance(topic.get("topics"), list):
        return topic["topics"]
    if isinstance(topic.get("children"), dict):
        return topic["children"].get("attached", []) or []
    return []


def _title_of(topic):
    return (topic.get("title") or "").strip()


def _node_to_text(topic):
    title = _title_of(topic)
    children = _children_of(topic)
    if not children:
        return title
    parts = [title] if title else []
    for c in children:
        parts.append(_node_to_text(c))
    return " / ".join([p for p in parts if p])


def _build_row_from_path(path_parts):
    row = {col: "" for col in COLUMNS}
    for i, val in enumerate(path_parts):
        if i in LEVEL_MAP:
            row[LEVEL_MAP[i]] = val
        else:
            break
    return row


def _extract_expected_actual_from_operation(topic):
    children = _children_of(topic)
    if not children:
        return "", ""
    expected = _node_to_text(children[0]) if len(children) >= 1 else ""
    actual = _node_to_text(children[1]) if len(children) >= 2 else ""
    return expected, actual


# ============================================================
# 方式一：xmindparser 解析
# ============================================================
def flatten_xmind(topic, path=None, records=None):
    if path is None:
        path = []
    if records is None:
        records = []

    title = _title_of(topic)
    current_path = path + [title] if title else path
    children = topic.get("topics", []) or []

    if len(current_path) == OPERATION_LEVEL:
        row = _build_row_from_path(current_path)
        exp, act = _extract_expected_actual_from_operation(topic)
        row["预期结果"] = exp
        row["实际结果"] = act
        row["完整路径"] = " > ".join(current_path + [exp, act])
        records.append(row)
        return records

    if not children:
        if current_path:
            row = _build_row_from_path(current_path)
            row["完整路径"] = " > ".join(current_path)
            records.append(row)
        return records

    for child in children:
        flatten_xmind(child, current_path, records)

    return records


def parse_xmind_file(uploaded_file):
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xmind") as tmp:
            tmp.write(uploaded_file.getvalue())
            tmp_path = tmp.name

        sheets = xmind_to_dict(tmp_path)
        all_records = []

        for sheet in sheets:
            root_topic = sheet.get("topic")
            if not root_topic:
                continue
            for child in root_topic.get("topics", []) or []:
                flatten_xmind(child, path=[], records=all_records)

        if not all_records:
            return None, "未在 XMind 中检测到符合规则的层级结构。"

        df = pd.DataFrame(all_records)
        df.insert(0, "序号", range(1, len(df) + 1))
        cols = ["序号"] + COLUMNS + ["完整路径"]
        df = df[[c for c in cols if c in df.columns]]
        return df, None

    except Exception as e:
        return None, f"解析失败：{e}"
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass


# ============================================================
# 方式二：zip + content.json + resources 截图
# ============================================================
def _collect_image_refs(topic):
    refs = set()
    img = topic.get("image")
    if isinstance(img, dict) and img.get("src"):
        refs.add(img["src"])

    for ext in topic.get("extensions", []) or []:
        if not isinstance(ext, dict):
            continue
        for att in ext.get("attachments", []) or []:
            if isinstance(att, dict) and att.get("path"):
                refs.add(att["path"])
        if ext.get("src"):
            refs.add(ext["src"])

    notes = topic.get("notes")
    if isinstance(notes, dict):
        plain = notes.get("plain", {}).get("content", "")
        if isinstance(plain, str):
            for m in re.findall(r"resources/[^\"'\)\s>]+", plain):
                refs.add(m)
    return refs


def _extract_images(topic, resource_images):
    imgs = []
    for ref in _collect_image_refs(topic):
        key = os.path.basename(ref.replace("\\", "/"))
        for cand in (ref, f"resources/{key}", key):
            if cand in resource_images:
                imgs.append(resource_images[cand])
                break
    return imgs


def _walk_content_json(topic, path, records, resource_images, inherited_imgs=None):
    if inherited_imgs is None:
        inherited_imgs = []

    title = _title_of(topic)
    current_path = path + [title] if title else path
    children = _children_of(topic)

    own_imgs = _extract_images(topic, resource_images)
    cur_imgs = own_imgs if own_imgs else inherited_imgs

    if len(current_path) == OPERATION_LEVEL:
        row = _build_row_from_path(current_path)
        exp, act = _extract_expected_actual_from_operation(topic)
        row["预期结果"] = exp
        row["实际结果"] = act

        child_imgs = cur_imgs
        for c in children:
            ci = _extract_images(c, resource_images)
            if ci:
                child_imgs = ci
                break

        row["完整路径"] = " > ".join(current_path + [exp, act])
        row["截图"] = list(child_imgs) if child_imgs else []
        records.append(row)
        return

    if not children:
        if current_path:
            row = _build_row_from_path(current_path)
            row["完整路径"] = " > ".join(current_path)
            row["截图"] = list(cur_imgs) if cur_imgs else []
            records.append(row)
        return

    for child in children:
        _walk_content_json(child, current_path, records, resource_images, cur_imgs)


def parse_xmind_zip(uploaded_file):
    try:
        data = uploaded_file.getvalue()
        zf = zipfile.ZipFile(io.BytesIO(data))
        names = zf.namelist()

        content_name = None
        for n in names:
            if n.endswith("content.json"):
                content_name = n
                break
        if not content_name:
            return None, "压缩包中未找到 content.json，可能不是 XMind Zen/2020+ 格式。"

        with zf.open(content_name) as f:
            content = json.loads(f.read().decode("utf-8"))

        resource_images = {}
        for n in names:
            norm = n.replace("\\", "/")
            if norm.startswith("resources/") or "/resources/" in f"/{norm}":
                if norm.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")):
                    b = zf.read(n)
                    resource_images[norm] = b
                    resource_images[os.path.basename(norm)] = b

        all_records = []
        sheets = content if isinstance(content, list) else [content]

        for sheet in sheets:
            root = sheet.get("rootTopic") or sheet.get("topic")
            if not root:
                continue
            for child in _children_of(root):
                _walk_content_json(child, [], all_records, resource_images)

        if not all_records:
            return None, "未在 content.json 中检测到符合规则的层级结构。"

        df = pd.DataFrame(all_records)
        df.insert(0, "序号", range(1, len(df) + 1))
        cols = ["序号"] + COLUMNS + ["完整路径", "截图"]
        df = df[[c for c in cols if c in df.columns]]
        return df, None

    except zipfile.BadZipFile:
        return None, "文件不是有效的 zip/xmind 包，无法解压。"
    except Exception as e:
        return None, f"解析失败：{e}"


# ============================================================
# 拆分
# ============================================================
def split_dataframes(df):
    if df is None or df.empty:
        empty = df.iloc[0:0] if df is not None else pd.DataFrame(columns=COLUMNS)
        return empty, empty, 0

    case_col = df["用例名称"].astype(str) if "用例名称" in df.columns else pd.Series([""] * len(df))
    module_col = df["模块名称"].astype(str) if "模块名称" in df.columns else pd.Series([""] * len(df))

    detail_mask = case_col.str.contains(DETAIL_KEYWORD, na=False)
    dropped_count = int(detail_mask.sum())
    kept = df[~detail_mask].copy()

    prepare_mask = kept["模块名称"].astype(str).str.contains(PREPARE_KEYWORD, na=False)
    prepare_df = kept[prepare_mask].copy().reset_index(drop=True)
    normal_df = kept[~prepare_mask].copy().reset_index(drop=True)

    for d in (normal_df, prepare_df):
        if "序号" in d.columns and not d.empty:
            d["序号"] = range(1, len(d) + 1)

    return normal_df, prepare_df, dropped_count


# ============================================================
# 导出 Excel（图片精确锚定到单元格）
# ============================================================
def _bytes_to_xl_image(img_bytes, max_width=IMG_MAX_WIDTH):
    """返回 (XLImage, w, h)；失败返回 (None, 0, 0)。"""
    try:
        bio = io.BytesIO(img_bytes)
        pil = PILImage.open(bio)
        if pil.mode not in ("RGB", "RGBA"):
            pil = pil.convert("RGB")
        w, h = pil.size
        if w > max_width:
            ratio = max_width / w
            w = int(w * ratio)
            h = int(h * ratio)
            pil = pil.resize((w, h))
        out = io.BytesIO()
        pil.save(out, format="PNG")
        out.seek(0)
        xl_img = XLImage(out)
        xl_img.width = w
        xl_img.height = h
        return xl_img, w, h
    except Exception:
        return None, 0, 0


def _add_image_to_cell(ws, img_bytes, row_idx, col_idx, offset_y_px=0):
    """
    把图片按 OneCellAnchor 方式插入到 (row_idx, col_idx) 单元格内。
    - row_idx、col_idx 均 1-based
    - offset_y_px：在该单元格内的纵向像素偏移（用于同一单元格堆叠多图）
    """
    xl_img, w, h = _bytes_to_xl_image(img_bytes)
    if xl_img is None:
        return 0

    # AnchorMarker 的 col/row 是 0-based
    marker = AnchorMarker(
        col=col_idx - 1,
        colOff=pixels_to_EMU(IMG_CELL_PADDING),
        row=row_idx - 1,
        rowOff=pixels_to_EMU(offset_y_px + IMG_CELL_PADDING),
    )
    size = XDRPositiveSize2D(
        cx=pixels_to_EMU(w),
        cy=pixels_to_EMU(h),
    )
    xl_img.anchor = OneCellAnchor(_from=marker, ext=size)
    ws.add_image(xl_img)
    return h


def _embed_images_into_row(ws, img_list, row_idx, img_col_idx):
    """在同一行的截图列内纵向堆叠多张图片，并按总高度调整行高。"""
    if not img_list:
        return

    col_letter = get_column_letter(img_col_idx)
    ws.column_dimensions[col_letter].width = IMG_COL_WIDTH

    total_height = 0
    for img_bytes in img_list:
        h = _add_image_to_cell(ws, img_bytes, row_idx, img_col_idx, offset_y_px=total_height)
        if h:
            total_height += h + IMG_CELL_PADDING

    if total_height > 0:
        # 像素 → 磅：1 磅 ≈ 1.333 像素，用 0.75 换算并留余量
        ws.row_dimensions[row_idx].height = max(
            IMG_ROW_HEIGHT,
            int(total_height * 0.75) + 10,
        )


def dataframe_to_excel_bytes(normal_df, prepare_df,
                             sheet_name="测试用例", embed_images=True):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:

        # ---------- Sheet 1：主用例 ----------
        safe_main = _safe_sheet_name(sheet_name)
        normal_write = normal_df.drop(columns=["截图"]) if "截图" in normal_df.columns else normal_df
        normal_write.to_excel(writer, index=False, sheet_name=safe_main)
        ws_main = writer.sheets[safe_main]

        if embed_images and "截图" in normal_df.columns:
            headers = list(normal_write.columns)
            img_col_idx = len(headers) + 1
            ws_main.cell(row=1, column=img_col_idx, value="截图")
            ws_main.column_dimensions[get_column_letter(img_col_idx)].width = IMG_COL_WIDTH

            for i, img_list in enumerate(normal_df["截图"].tolist()):
                if not img_list:
                    continue
                _embed_images_into_row(ws_main, img_list, row_idx=2 + i, img_col_idx=img_col_idx)

        # ---------- Sheet 2：需求详情（放测试准备） ----------
        safe_detail = _safe_sheet_name(DETAIL_SHEET_NAME)

        if prepare_df is not None and not prepare_df.empty:
            prepare_write = prepare_df.drop(columns=["截图"]) if "截图" in prepare_df.columns else prepare_df
            prepare_write.to_excel(writer, index=False, sheet_name=safe_detail, startrow=1)
            ws_detail = writer.sheets[safe_detail]

            if embed_images and "截图" in prepare_df.columns:
                headers = list(prepare_write.columns)
                img_col_idx = len(headers) + 1
                ws_detail.cell(row=2, column=img_col_idx, value="截图")
                ws_detail.column_dimensions[get_column_letter(img_col_idx)].width = IMG_COL_WIDTH

                for i, img_list in enumerate(prepare_df["截图"].tolist()):
                    if not img_list:
                        continue
                    _embed_images_into_row(ws_detail, img_list, row_idx=3 + i, img_col_idx=img_col_idx)
        else:
            pd.DataFrame().to_excel(writer, index=False, sheet_name=safe_detail)

    return output.getvalue()


# ============================================================
# 使用说明（主页面折叠面板）
# ============================================================
with st.expander("📖 使用说明与规则（点击展开）", expanded=False):
    col_a, col_b = st.columns(2)

    with col_a:
        st.markdown("#### 📌 使用步骤")
        st.markdown(
            """
            1. 上传 `.xmind` 文件
            2. 选择解析方式：
               - **常规解析**：xmindparser（快速，不含截图）
               - **解压解析（含截图）**：读 `content.json` + `resources/` 截图
            3. 预览后下载 Excel
            """
        )

        st.markdown("#### 📐 节点与表头对应")
        st.markdown(
            """
            - 一级节点 → **模块名称**
            - 二级节点 → **用例名称**
            - 三级节点 → **操作步骤**
            - 四级第 1 个 → **预期结果**
            - 四级第 2 个 → **实际结果**
            """
        )

    with col_b:
        st.markdown("#### 🗂️ Sheet 规则")
        st.markdown(
            f"""
            - **Sheet1**（XMind 文件名）：普通用例
            - **Sheet2「{DETAIL_SHEET_NAME}」**：
                - **A2 起** = 模块名称含「{PREPARE_KEYWORD}」的表格
            """
        )

        st.markdown("#### 🖼️ 图片规则")
        st.markdown(
            """
            - 图片插入到**对应用例行**「实际结果」列之后
            - 多张图片在同一单元格内纵向堆叠
            - 行高按图片高度自动调整
            """
        )

        st.markdown("#### 🚫 过滤规则")
        st.markdown(
            f"""
            - 用例名称（二级节点）含「**{DETAIL_KEYWORD}**」→ **不导出**
            """
        )


# ============================================================
# 主界面
# ============================================================
uploaded_file = st.file_uploader(
    "选择 XMind 文件",
    type=["xmind"],
    help="支持 XMind 8 及 XMind Zen / 2020+ 格式",
)

if uploaded_file is not None:
    xmind_base = os.path.splitext(uploaded_file.name)[0]
    excel_name = f"{xmind_base}.xlsx"
    main_sheet_name = _safe_sheet_name(xmind_base)

    col1, col2 = st.columns(2)
    with col1:
        btn_normal = st.button("🚀 常规解析（xmindparser）", use_container_width=True)
    with col2:
        btn_zip = st.button("📦 解压解析（含截图）", use_container_width=True)

    if btn_normal:
        with st.spinner("正在解析 XMind 文件..."):
            df, error = parse_xmind_file(uploaded_file)
        st.session_state["xmind_result"] = (df, error, "normal")

    if btn_zip:
        with st.spinner("正在解压并解析 content.json..."):
            df, error = parse_xmind_zip(uploaded_file)
        st.session_state["xmind_result"] = (df, error, "zip")

    with st.expander("🔍 调试：查看 content.json 结构", expanded=False):
        try:
            zf = zipfile.ZipFile(io.BytesIO(uploaded_file.getvalue()))
            found = False
            for n in zf.namelist():
                if n.endswith("content.json"):
                    c = json.loads(zf.read(n).decode("utf-8"))
                    st.json(c, expanded=False)
                    found = True
                    break
            if not found:
                st.warning("未找到 content.json")
        except Exception as e:
            st.warning(f"读取 content.json 失败：{e}")

    if "xmind_result" in st.session_state:
        df, error, mode = st.session_state["xmind_result"]

        if error:
            st.error(error)
        else:
            normal_df, prepare_df, dropped_count = split_dataframes(df)

            has_img = (
                ("截图" in df.columns and df["截图"].apply(lambda x: bool(x)).any())
                if df is not None and not df.empty else False
            )

            st.success(
                f"✅ 解析成功：主用例 **{len(normal_df)}** 条，"
                f"测试准备 **{len(prepare_df)}** 条，"
                f"已过滤需求详情 **{dropped_count}** 条"
                + ("，含截图。" if has_img else "。")
            )
            st.caption(
                f"导出文件名：`{excel_name}`　|　"
                f"Sheet1：`{main_sheet_name}`　|　"
                f"Sheet2：`{DETAIL_SHEET_NAME}`"
            )

            st.subheader(f"数据预览 - 主用例（{len(normal_df)} 条）")
            pn = normal_df.copy()
            if "截图" in pn.columns:
                pn["截图"] = pn["截图"].apply(
                    lambda lst: f"🖼️ {len(lst)} 张" if isinstance(lst, list) and lst else ""
                )
            st.dataframe(pn, use_container_width=True, height=400)

            if not prepare_df.empty:
                st.subheader(f"数据预览 - 测试准备（{len(prepare_df)} 条）")
                pp = prepare_df.copy()
                if "截图" in pp.columns:
                    pp["截图"] = pp["截图"].apply(
                        lambda lst: f"🖼️ {len(lst)} 张" if isinstance(lst, list) and lst else ""
                    )
                st.dataframe(pp, use_container_width=True, height=200)

            if has_img:
                with st.expander("🖼️ 查看用例截图", expanded=False):
                    for _, row in df.iterrows():
                        imgs = row.get("截图")
                        if isinstance(imgs, list) and imgs:
                            st.markdown(f"**{row.get('完整路径', '')}**")
                            for b in imgs:
                                st.image(b, width=320)

            excel_bytes = dataframe_to_excel_bytes(
                normal_df, prepare_df,
                sheet_name=main_sheet_name,
                embed_images=has_img,
            )

            st.download_button(
                label="⬇️ 下载 Excel 文件",
                data=excel_bytes,
                file_name=excel_name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
else:
    st.info("请上传一个 `.xmind` 文件开始解析。")